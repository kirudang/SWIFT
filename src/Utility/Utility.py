import argparse
import json
import time
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

from openai import OpenAI
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.formatting.rule import ColorScaleRule

# -----------------------------------------------------------------------------
# API OPENAI 
# -----------------------------------------------------------------------------
OPENAI_API_KEY = "Your OpenAI API Key"
MODEL = "gpt-5.5"

# Reference / Original
DATA_PATH_A = "./Data/C4_KGW_output.json"
FIELD_A = "Original_summary"

# Testing / Watermarked
DATA_PATH_B = DATA_PATH_A = "./Data/C4_KGW_output.json"
FIELD_B = "Watermarked_summary"

OUTPUT_PATH = "KGW_Gemma_C4.xlsx"

MAX_SAMPLES = 1000
WORKERS = 10


# ============================================================
# JUDGE PROMPT
# ============================================================

SYSTEM_PROMPT = """
You are an expert evaluator of semantic and factual consistency.

Given an original text (Text A) and a watermarked text (Text B),
evaluate how well Text B maintains the core meaning and factual
content represented in Text A.

Text B does not need to reproduce the exact wording, sentence structure,
or level of detail of Text A. Different but semantically compatible
expressions or formulations should not be penalized.

Consider:
1. Overall semantic meaning and consistency.
2. Named entities, including people, organizations, locations,
   products, titles, and technical terms.
3. Whether important information is contradicted, omitted, or
   materially changed.
4. Numbers, dates, quantities, units, relations, actions, roles,
   and polarity.

Scoring rubric:
0 = Completely unrelated, contradictory, or the core meaning is not preserved.
1 = Same general topic, but most important meaning or facts differ.
2 = Partial semantic consistency, with substantial important information
    changed, contradicted, or missing.
3 = Mostly maintains the core meaning, but some important details,
    entities, or facts differ.
4 = Maintains the core meaning and factual content, with only minor
    semantic or factual differences.
5 = Semantically and factually consistent; key entities, numbers,
    relations, and important information are maintained, even if the
    wording, structure, or presentation differs.

Important:
- If Text B directly contradicts important information in Text A,
  assign a score of at most 1.
- Do not penalize Text B simply for using different wording, sentence
  structure, phrasing, or level of detail.
- Do not require lexical or stylistic similarity between the two texts.
- Focus on substantive semantic and factual differences rather than
  surface-form differences.
- Treat Text A as a reference for semantic and factual consistency,
  not as a template that Text B must closely reproduce.

"""


# ============================================================
# LOAD DATA
# ============================================================

def load_json(path):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def load_pairs():
    rows_a = load_json(DATA_PATH_A)

    if Path(DATA_PATH_A).resolve() == Path(DATA_PATH_B).resolve():
        rows_b = rows_a
    else:
        rows_b = load_json(DATA_PATH_B)

    n = min(len(rows_a), len(rows_b), MAX_SAMPLES)

    pairs = []

    for i in range(n):
        if FIELD_A not in rows_a[i]:
            raise KeyError(f"Row {i} missing {FIELD_A} in DATA_PATH_A")

        if FIELD_B not in rows_b[i]:
            raise KeyError(f"Row {i} missing {FIELD_B} in DATA_PATH_B")

        pairs.append(
            (
                str(rows_a[i][FIELD_A] or ""),
                str(rows_b[i][FIELD_B] or ""),
            )
        )

    return pairs


# ============================================================
# OPENAI JUDGE
# ============================================================

def judge_one(client, index, text_a, text_b):
    prompt = f"""
Text A (Original / Reference):
{text_a}

Text B (Watermarked / Compared):
{text_b}

Return JSON only:
{{
    "score": 0,
    "reason": "brief explanation"
}}
"""

    for attempt in range(5):
        try:
            response = client.responses.create(
                model=MODEL,
                instructions=SYSTEM_PROMPT,
                input=prompt,
                text={
                    "format": {
                        "type": "json_schema",
                        "name": "semantic_similarity",
                        "schema": {
                            "type": "object",
                            "properties": {
                                "score": {
                                    "type": "integer",
                                    "minimum": 0,
                                    "maximum": 5,
                                },
                                "reason": {
                                    "type": "string",
                                },
                            },
                            "required": ["score", "reason"],
                            "additionalProperties": False,
                        },
                        "strict": True,
                    }
                },
            )

            result = json.loads(response.output_text)

            return {
                "index": index,
                "score": int(result["score"]),
                "reason": result["reason"],
                "error": "",
            }

        except Exception as e:
            if attempt == 4:
                return {
                    "index": index,
                    "score": None,
                    "reason": "",
                    "error": str(e),
                }

            time.sleep(2 ** attempt)


# ============================================================
# EVALUATE
# ============================================================

def evaluate(pairs):
    client = OpenAI(api_key=OPENAI_API_KEY)

    results = [None] * len(pairs)

    with ThreadPoolExecutor(max_workers=WORKERS) as executor:
        futures = {
            executor.submit(
                judge_one,
                client,
                i,
                text_a,
                text_b,
            ): i
            for i, (text_a, text_b) in enumerate(pairs)
        }

        completed = 0

        for future in as_completed(futures):
            i = futures[future]
            results[i] = future.result()

            completed += 1

            if completed % 10 == 0 or completed == len(pairs):
                print(f"Completed {completed}/{len(pairs)}")

    return results


# ============================================================
# SAVE EXCEL
# ============================================================

def save_excel(pairs, results):
    wb = Workbook()

    ws = wb.active
    ws.title = "Semantic Scores"

    ws.append([
        "Data Point",
        FIELD_A,
        FIELD_B,
        "Semantic Score (0-5)",
        "Evaluation Note",
        "Error",
    ])

    for i, ((text_a, text_b), result) in enumerate(
        zip(pairs, results),
        start=1,
    ):
        ws.append([
            i,
            text_a,
            text_b,
            result["score"],
            result["reason"],
            result["error"],
        ])

    ws.freeze_panes = "A2"

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 60
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 60
    ws.column_dimensions["F"].width = 30

    # Summary
    scores = [
        r["score"]
        for r in results
        if r["score"] is not None
    ]

    distribution = Counter(scores)
    average = sum(scores) / len(scores) if scores else 0

    summary = wb.create_sheet("Summary")

    summary.append(["Metric", "Value"])
    summary.append(["Model", MODEL])
    summary.append(["Reference field", FIELD_A])
    summary.append(["Compared field", FIELD_B])
    summary.append(["Total samples", len(pairs)])
    summary.append(["Successfully scored", len(scores)])
    summary.append(["Average score", average])

    for score in range(5, -1, -1):
        summary.append([
            f"Score {score}",
            distribution.get(score, 0),
        ])

    wb.save(OUTPUT_PATH)

    print("\nFinished")
    print(f"Average score: {average:.3f} / 5.000")
    print(f"Distribution: {dict(sorted(distribution.items(), reverse=True))}")
    print(f"Saved: {OUTPUT_PATH}")


# ============================================================
# MAIN
# ============================================================

if __name__ == "__main__":

    if not OPENAI_API_KEY:
        raise ValueError("Add your OpenAI API key to OPENAI_API_KEY.")

    pairs = load_pairs()

    print(f"Model: {MODEL}")
    print(f"Reference: {DATA_PATH_A} -> {FIELD_A}")
    print(f"Compared:  {DATA_PATH_B} -> {FIELD_B}")
    print(f"Samples:   {len(pairs)}")
    print()

    results = evaluate(pairs)

    save_excel(pairs, results)
